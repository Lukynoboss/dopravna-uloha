from __future__ import annotations

import json
import sys
import time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from statistiky import HIST_HRANY, VZORKA_ULOH, Statistiky

KOREN = Path(__file__).resolve().parent

NAZVY = {"SZR": "Metóda severozápadného rohu", "IM": "Indexová metóda", "VAM": "Vogelova aproximačná metóda",
         "JVR": "Metóda juhovýchodného rohu (nová)", "MDD": "Metóda dvojitej diferencie (nová)",
         "PVM": "Predvídavá metóda (nová)"}
FARBY = {"SZR": "2B7BD6", "IM": "EE6A2C", "VAM": "1FAE7A", "JVR": "8E5BD0",
         "MDD": "0F8B8D", "PVM": "D4A017"}
ROHY = {"SZR", "JVR"}
POJMY_METOD = {
    "SZR": "Plní tabuľku od ľavého horného rohu („severozápad“) doprava a dole, vôbec nehľadí na ceny. Je "
           "najjednoduchšia, ale výsledok býva drahý.",
    "IM": "V každom kroku vyberie najlacnejšie voľné pole tabuľky a pošle cezeň čo najviac. Hľadí na ceny, preto je "
          "oveľa lepšia ako SZR.",
    "VAM": "Pre každý riadok a stĺpec spočíta rozdiel dvoch najnižších cien („pokuta“, ak nevyberieme najlacnejšie "
           "pole). Najprv obslúži riadok alebo stĺpec s najväčšou pokutou, a to jeho najlacnejším poľom. Je "
           "náročnejšia na výpočet, ale dáva riešenie najbližšie k optimu.",
    "JVR": "Nová (vymyslená) metóda: ako SZR, ale začína v pravom dolnom rohu. Dáva vždy ten istý plán ako SZR.",
    "MDD": "Nová (vymyslená) metóda dvojitej diferencie: pole dostane diferenciu riadku, ak je v ňom najlacnejšie, "
           "plus diferenciu stĺpca, ak je najlacnejšie v stĺpci; obsadí sa pole s najväčším súčtom. Kvalitou blízko VAM.",
    "PVM": "Nová (vymyslená) predvídavá metóda: v každom kroku vyskúša 3 najlepšie polia podľa MDD, zvyšok dopočíta "
           "VAM a vyberie pole s najnižšou celkovou cenou. Výrazne lepšia ako VAM, ale počíta viac.",
}
TYPY = {"vyrovnana": "vyrovnaná", "prebytok_ponuky": "prebytok ponuky", "prebytok_dopytu": "prebytok dopytu"}

PISMO = "Arial"
F_ZAKLAD = Font(name=PISMO, size=10)
F_TUCNE = Font(name=PISMO, size=10, bold=True)
F_NADPIS = Font(name=PISMO, size=16, bold=True, color="1F2A3A")
F_PODNADPIS = Font(name=PISMO, size=12, bold=True, color="1F4E79")
F_SIVE = Font(name=PISMO, size=9, italic=True, color="5F6B7A")
F_HLAVICKA = Font(name=PISMO, size=10, bold=True, color="FFFFFF")
FILL_HLAVICKA = PatternFill("solid", fgColor="1F4E79")
FILL_VYSVETLIVKA = PatternFill("solid", fgColor="EEF3F8")
FILL_ZVYRAZNENIE = PatternFill("solid", fgColor="E7F5EE")
OKRAJ = Border(*(Side(style="thin", color="C9D2DC"),) * 4)
ZALOM = Alignment(wrap_text=True, vertical="top")
STRED = Alignment(horizontal="center", vertical="center", wrap_text=True)

PCT = "0.0%"
PCT2 = "0.00%"
DES1 = "0.0"
DES2 = "0.00"
CELE = "#,##0"


def cislo(x: float, des: int = 1) -> str:
    return f"{x:,.{des}f}".replace(",", " ").replace(".", ",")


def bunka(ws, ref, hodnota, font=F_ZAKLAD, fmt=None, fill=None, align=None, okraj=False):
    c = ws[ref]
    c.value = hodnota
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if okraj:
        c.border = OKRAJ
    return c


def odsek(ws, riadok: int, text: str, stlpce: int = 8, font=F_ZAKLAD, vyska: float | None = None):
    ws.merge_cells(start_row=riadok, start_column=1, end_row=riadok, end_column=stlpce)
    c = ws.cell(riadok, 1, text)
    c.font = font
    c.alignment = ZALOM
    if vyska is None:
        sirka = sum((ws.column_dimensions[get_column_letter(i)].width or 10) for i in range(1, stlpce + 1))
        riadkov = max(1, sum(len(cast) // max(1, int(sirka * 1.1)) + 1 for cast in text.split("\n")))
        vyska = 14 * riadkov + 2
    ws.row_dimensions[riadok].height = vyska
    return riadok + 1


def hlavicka(ws, riadok: int, nazvy: list[str], vysvetlivky: list[str] | None = None, stlpec: int = 1):
    for i, n in enumerate(nazvy):
        bunka(ws, ws.cell(riadok, stlpec + i).coordinate, n, F_HLAVICKA, fill=FILL_HLAVICKA, align=STRED, okraj=True)
    ws.row_dimensions[riadok].height = 44
    if vysvetlivky:
        for i, v in enumerate(vysvetlivky):
            bunka(ws, ws.cell(riadok + 1, stlpec + i).coordinate, v, F_SIVE, fill=FILL_VYSVETLIVKA,
                  align=ZALOM, okraj=True)
        ws.row_dimensions[riadok + 1].height = max(30, 13 * (max(len(v) for v in vysvetlivky) // 13 + 1))
        return riadok + 2
    return riadok + 1


def sirky(ws, hodnoty: list[float]):
    for i, w in enumerate(hodnoty, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def osi(graf):
    graf.x_axis.delete = False
    graf.y_axis.delete = False


def farby_stlpcov(graf, M):
    for i, m in enumerate(M):
        bod = DataPoint(idx=i)
        bod.graphicalProperties.solidFill = FARBY[m]
        bod.graphicalProperties.line.solidFill = FARBY[m]
        graf.series[0].dPt.append(bod)


def zistenia(s: Statistiky) -> list[str]:
    st = s.stat
    M = s.metody
    k = {m: max(1, st[m][0]) for m in M}
    odch = {m: st[m][2] / k[m] for m in M}
    opt = {m: 100 * st[m][1] / k[m] for m in M}
    kroky = {m: st[m][3] / k[m] for m in M}
    cas = {m: st[m][4] / k[m] / 1e6 for m in M}
    najl = {m: 100 * s.najlepsia[m] / max(1, s.hotovo) for m in M}
    poradie = sorted(M, key=odch.get)
    b, t = poradie[0], poradie[-1]
    ostatne = poradie[1:]
    vety = [
        f"1. Najsprávnejšia je {NAZVY[b]} ({b}). Jej prvé (východiskové) riešenie je v priemere len o "
        f"{cislo(odch[b])} % drahšie ako najlacnejšie možné riešenie (optimum)."
        + (" Ostatné: " + ", ".join(f"{m} {cislo(odch[m])} %" for m in ostatne) + "." if ostatne else ""),
        f"2. {b} trafila optimum hneď na prvý pokus v {cislo(opt[b], 2)} % úloh"
        + (" (" + ", ".join(f"{m}: {cislo(opt[m], 2)} %" for m in ostatne) + ")" if ostatne else "")
        + f". Keď porovnáme všetky metódy na tej istej úlohe, {b} dala najlacnejšie východiskové riešenie "
        f"v {cislo(najl[b])} % úloh.",
    ]
    pk = sorted(M, key=kroky.get)
    vety.append(
        f"3. Najmenej práce po východiskovom riešení zostane pri {pk[0]}: metóda MODI ho musí opraviť priemerne "
        f"{cislo(kroky[pk[0]])}-krát"
        + (" (" + ", ".join(f"{m}: {cislo(kroky[m])}" for m in pk[1:]) + " krokov)" if len(pk) > 1 else "")
        + ". Jeden krok MODI = jedno prerozdelenie prepravy po uzavretom obvode v tabuľke.")
    pc = sorted(M, key=cas.get)
    vety.append(
        f"4. V celkovom čase výpočtu (metóda + MODI) je najrýchlejšia {pc[0]} ({cislo(cas[pc[0]], 2)} ms na úlohu)"
        + (", potom " + ", ".join(f"{m} ({cislo(cas[m], 2)} ms)" for m in pc[1:]) if len(pc) > 1 else "")
        + ". Čas závisí od počítača a programu, preto je spoľahlivejším meradlom práce počet krokov MODI.")
    vel = sorted(s.velkost)
    if len(vel) >= 2:
        mala, velka = vel[0], vel[-1]
        g1, g2 = s.velkost[mala][b], s.velkost[velka][b]
        vety.append(
            f"5. Čím je úloha väčšia, tým sú všetky metódy horšie. Pri najmenších úlohách (väčší rozmer {mala}) má "
            f"{b} priemernú odchýlku {cislo(g1[1] / max(1, g1[0]))} % a trafí optimum v "
            f"{cislo(100 * g1[3] / max(1, g1[0]))} % úloh; pri najväčších (väčší rozmer {velka}) je odchýlka "
            f"{cislo(g2[1] / max(1, g2[0]))} % a optimum trafí v {cislo(100 * g2[3] / max(1, g2[0]))} % úloh.")
    if len(s.typ) >= 2 and "vyrovnana" in s.typ:
        ost = [t_ for t_ in s.typ if t_ != "vyrovnana"]
        gv = s.typ["vyrovnana"][b]
        texty = ", ".join(f"pri type „{TYPY.get(t_, t_)}“ {cislo(s.typ[t_][b][1] / max(1, s.typ[t_][b][0]))} %"
                          for t_ in ost)
        vety.append(
            f"6. Nevyrovnané úlohy (ponuka ≠ dopyt, treba pridať fiktívneho dodávateľa alebo odberateľa) sú ťažšie: "
            f"odchýlka {b} je pri vyrovnaných úlohách {cislo(gv[1] / max(1, gv[0]))} %, {texty}.")
    zaver = f"Záver: {b} je najlepšia voľba – dá riešenie najbližšie k optimu."
    if t != b:
        zaver += f" Najďalej od optima býva {t}" + (", lebo vôbec nehľadí na ceny." if t in ROHY else ".")
    vety.append(zaver)
    return vety


def exportuj(cesta: Path, s: Statistiky, info: dict) -> Path:
    cesta = Path(cesta)
    wb = Workbook()
    p = info.get("parametre_generatora") or {}
    M = s.metody
    k_ = len(M)
    zoznam = ", ".join(M[:-1]) + (" a " if len(M) > 1 else "") + M[-1]

    wp = wb.active
    wp.title = "Podklady"
    sirky(wp, [12, 12, 16, 18, 18, 20, 20, 16, 18])
    r = 1
    bunka(wp, "A1", "Podklady – súčty priamo zo simulácie", F_PODNADPIS)
    r = odsek(wp, 2, "Tieto čísla zapísal program počas výpočtu. Ostatné hárky z nich počítajú priemery a percentá "
                     "vzorcami (napr. priemerná odchýlka = súčet odchýlok / počet úloh). Tu nič nemeň.", 9, F_SIVE)
    r = hlavicka(wp, 4, ["metóda", "počet úloh", "trafila optimum (počet úloh)", "súčet odchýlok (v %)",
                         "súčet krokov MODI", "súčet času metódy (ms)", "súčet času MODI (ms)",
                         "najlepšia v úlohe (počet)", "degenerované riešenie (počet)"])
    PODKLADY_RIADOK = {}
    for m in M:
        k, opt, odch, kroky, _, tmet, tmodi, deg = s.stat[m]
        hodnoty = [m, k, opt, odch, kroky, tmet / 1e6, tmodi / 1e6, s.najlepsia[m], deg]
        for i, h in enumerate(hodnoty):
            bunka(wp, wp.cell(r, i + 1).coordinate, h, F_TUCNE if i == 0 else F_ZAKLAD,
                  CELE if i in (1, 2, 4, 7, 8) else DES2, okraj=True)
        PODKLADY_RIADOK[m] = r
        r += 1
    r += 1
    bunka(wp, f"A{r}", "Počet úloh spolu", F_TUCNE)
    bunka(wp, f"B{r}", s.hotovo, fmt=CELE)
    bunka(wp, f"A{r + 1}", "Chyby počas výpočtu", F_TUCNE)
    bunka(wp, f"B{r + 1}", s.chyby, fmt=CELE)

    wv = wb.create_sheet("Výsledky metód", 0)
    sirky(wv, [9, 37, 12, 14, 14, 16, 14, 14, 13, 13, 13, 15, 11])
    bunka(wv, "A1", "Výsledky metód – priemer cez všetky úlohy", F_PODNADPIS)
    odsek(wv, 2, "Každý riadok je jedna metóda. Pod názvom stĺpca je napísané, čo číslo znamená. Hodnoty sú "
                 "vzorce z hárku „Podklady“ – keď klikneš na bunku, uvidíš, ako sa počíta. Zelený riadok = "
                 "metóda s najmenšou priemernou odchýlkou.", 13, F_SIVE)
    stlpce = ["skratka", "metóda", "počet úloh", "trafila optimum", "priemerná odchýlka od optima",
              "kroky MODI (priemer)", "najlepšia v úlohe", "čas metódy (ms)", "čas MODI (ms)",
              "celkový čas (ms)", "degenerované riešenie", "poradie podľa odchýlky"]
    vysv = ["", "plný názov metódy", "koľko úloh táto metóda riešila",
            "v koľkých % úloh bolo prvé riešenie hneď najlacnejšie možné",
            "o koľko % je prvé riešenie v priemere drahšie ako optimum (menej = lepšie)",
            "koľkokrát musela MODI riešenie opraviť, kým našla optimum (menej = menej práce)",
            "v koľkých % úloh dala táto metóda najlacnejšie zo všetkých riešení (pri zhode sa počíta každej)",
            "priemerný čas samotnej metódy na jednu úlohu",
            "priemerný čas dopočítania optima metódou MODI",
            "metóda + MODI spolu; závisí od počítača",
            "v koľkých % úloh malo prvé riešenie menej obsadených polí, než treba (doplnili sa nulou)",
            "1 = najlepšia (najmenšia priemerná odchýlka)"]
    r0 = hlavicka(wv, 4, stlpce, vysv)
    prvy = r0
    for i, m in enumerate(M):
        r = r0 + i
        pr = PODKLADY_RIADOK[m]
        hodnoty = [
            (m, None), (NAZVY[m], None), (f"=Podklady!B{pr}", CELE),
            (f"=IFERROR(Podklady!C{pr}/Podklady!B{pr},0)", PCT2),
            (f"=IFERROR(Podklady!D{pr}/Podklady!B{pr}/100,0)", PCT),
            (f"=IFERROR(Podklady!E{pr}/Podklady!B{pr},0)", DES1),
            (f"=IFERROR(Podklady!H{pr}/Podklady!B{pr},0)", PCT),
            (f"=IFERROR(Podklady!F{pr}/Podklady!B{pr},0)", "0.000"),
            (f"=IFERROR(Podklady!G{pr}/Podklady!B{pr},0)", "0.000"),
            (f"=H{r}+I{r}", "0.000"),
            (f"=IFERROR(Podklady!I{pr}/Podklady!B{pr},0)", PCT),
            (f"=RANK(E{r},$E${prvy}:$E${prvy + k_ - 1},1)", "0"),
        ]
        for j, (h, fmt) in enumerate(hodnoty):
            c = bunka(wv, wv.cell(r, j + 1).coordinate, h, F_ZAKLAD, fmt, okraj=True,
                      align=Alignment(horizontal="left" if j < 2 else "right"))
            if j == 0:
                c.font = Font(name=PISMO, size=10, bold=True, color=FARBY[m])
    posledny = r0 + k_ - 1
    najlepsia = min(M, key=lambda m: s.stat[m][2] / max(1, s.stat[m][0]))
    for j in range(1, len(stlpce) + 1):
        wv.cell(r0 + M.index(najlepsia), j).fill = FILL_ZVYRAZNENIE
    wv.freeze_panes = wv.cell(r0, 3)

    graf = BarChart()
    graf.type = "col"
    graf.title = "Priemerná odchýlka prvého riešenia od optima"
    graf.y_axis.numFmt = "0%"
    graf.y_axis.majorGridlines = None
    graf.add_data(Reference(wv, min_col=5, min_row=r0, max_row=posledny), titles_from_data=False)
    graf.set_categories(Reference(wv, min_col=1, min_row=r0, max_row=posledny))
    graf.legend = None
    osi(graf)
    farby_stlpcov(graf, M)
    graf.height, graf.width = 7.5, 13
    wv.add_chart(graf, f"A{posledny + 3}")
    graf2 = BarChart()
    graf2.type = "col"
    graf2.title = "Priemerný počet krokov MODI do optima"
    graf2.y_axis.majorGridlines = None
    graf2.add_data(Reference(wv, min_col=6, min_row=r0, max_row=posledny), titles_from_data=False)
    graf2.set_categories(Reference(wv, min_col=1, min_row=r0, max_row=posledny))
    graf2.legend = None
    osi(graf2)
    farby_stlpcov(graf2, M)
    graf2.height, graf2.width = 7.5, 13
    wv.add_chart(graf2, f"G{posledny + 3}")

    wo = wb.create_sheet("Prehľad", 0)
    sirky(wo, [30, 22, 60, 12, 12, 12, 12, 12])
    bunka(wo, "A1", "Porovnanie metód východiskového riešenia dopravnej úlohy", F_NADPIS)
    bunka(wo, "A2", f"Výsledky simulácie · {info.get('datum', time.strftime('%Y-%m-%d %H:%M'))}", F_SIVE)
    r = 4
    bunka(wo, f"A{r}", "Čo sme robili", F_PODNADPIS)
    r = odsek(wo, r + 1,
              f"Počítač vygeneroval {cislo(s.hotovo, 0)} náhodných dopravných úloh. Každá úloha je tabuľka: "
              "niekoľko dodávateľov (koľko tovaru majú), niekoľko odberateľov (koľko tovaru potrebujú) a cena "
              "prepravy jednej jednotky z každého dodávateľa ku každému odberateľovi. Cieľom je rozvrhnúť prepravu "
              "tak, aby bola celková cena čo najnižšia.\n"
              f"Každú úlohu sme vyriešili metódami, ktoré nájdu prvé (východiskové) riešenie – {zoznam}. "
              "Potom sme z každého východiskového riešenia pokračovali metódou MODI, až kým sme nenašli "
              "najlacnejšie možné riešenie (optimum). Všetky metódy riešili vždy tie isté úlohy, takže "
              "porovnanie je spravodlivé.\n"
              "Pýtame sa: ktorá metóda dá hneď na začiatku riešenie najbližšie k optimu (správnosť) a koľko "
              "práce potom ešte zostane – počet krokov MODI a čas (efektívnosť).", 3)
    r += 1
    bunka(wo, f"A{r}", "Hlavné zistenia", F_PODNADPIS)
    r += 1
    for veta in zistenia(s):
        r = odsek(wo, r, veta, 3, F_TUCNE if veta.startswith("Záver") else F_ZAKLAD)
    r += 1
    bunka(wo, f"A{r}", "Nastavenie simulácie", F_PODNADPIS)
    r = hlavicka(wo, r + 1, ["nastavenie", "hodnota", "čo to znamená"])
    ceny = p.get("ceny", {})
    typ = p.get("typ", {})
    nastavenia = [
        ("Počet vypočítaných úloh", s.hotovo, f"koľko náhodných úloh sa vyriešilo (každá metódami {zoznam})"),
        ("Metódy", zoznam, "ktoré metódy východiskového riešenia sa porovnávali"),
        ("Zadaný počet úloh", info.get("pocet_zadany", s.hotovo),
         "ak je väčší než vypočítaný, výpočet sa zastavil skôr"),
        ("Dodávatelia m", "–".join(map(str, p.get("m_rozsah", []))), "počet riadkov tabuľky, vyberá sa náhodne"),
        ("Odberatelia n", "–".join(map(str, p.get("n_rozsah", []))), "počet stĺpcov tabuľky, vyberá sa náhodne"),
        ("Ceny prepravy c_ij", ", ".join(f"{k_} ({cislo(100 * v, 0)} % úloh)" for k_, v in ceny.items()),
         "z akého rozsahu sa losujú ceny za prepravu jednej jednotky"),
        ("Typy úloh", ", ".join(f"{TYPY.get(k_, k_)} {cislo(100 * v, 0)} %" for k_, v in typ.items()),
         "vyrovnaná = ponuka sa rovná dopytu; inak sa pridá fiktívny dodávateľ alebo odberateľ s cenou 0"),
        ("Kapacity a požiadavky", "–".join(map(str, p.get("kapacity_zaklad", []))),
         "koľko tovaru majú dodávatelia / chcú odberatelia (základ, celé čísla)"),
        ("Seed", info.get("seed", ""), "„semienko“ náhody – s rovnakým seedom vzniknú presne tie isté úlohy"),
        ("Procesy (jadrá)", info.get("procesy", ""), "koľko jadier procesora počítalo naraz"),
        ("Čas výpočtu", f"{cislo(info.get('trvanie_s', 0), 1)} s", "ako dlho trval celý výpočet"),
        ("Chyby", s.chyby, "počet úloh, pri ktorých nastala chyba (má byť 0)"),
    ]
    for nazov, hodnota, popis in nastavenia:
        bunka(wo, f"A{r}", nazov, F_TUCNE, okraj=True)
        bunka(wo, f"B{r}", hodnota, fmt=CELE if isinstance(hodnota, int) else None, okraj=True,
              align=Alignment(horizontal="left", wrap_text=True, vertical="top"))
        bunka(wo, f"C{r}", popis, F_SIVE, align=ZALOM, okraj=True)
        r += 1
    r += 1
    bunka(wo, f"A{r}", "Čo je v jednotlivých hárkoch", F_PODNADPIS)
    r += 1
    harky = [
        ("Výsledky metód", "hlavná tabuľka: správnosť a efektívnosť každej metódy + 2 grafy"),
        ("Podľa veľkosti", "ako sa metódy správajú pri malých a veľkých úlohách + grafy"),
        ("Podľa typu úlohy", "vyrovnané vs. nevyrovnané úlohy"),
        ("Rozdelenie odchýlky", "koľko úloh malo malú, strednú a veľkú odchýlku od optima"),
        ("Údaje (vzorka)", f"riadky zo simulácie – prvých {cislo(min(s.hotovo, VZORKA_ULOH), 0)} úloh, dá sa filtrovať"),
        ("Slovník pojmov", "vysvetlenie skratiek a pojmov + ako to vysvetliť vyučujúcemu"),
        ("Podklady", "súčty zo simulácie, z ktorých ostatné hárky počítajú vzorcami"),
    ]
    for nazov, popis in harky:
        bunka(wo, f"A{r}", nazov, F_TUCNE)
        wo.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        bunka(wo, f"B{r}", popis)
        r += 1

    wg = wb.create_sheet("Podľa veľkosti")
    opt_m = "VAM" if "VAM" in M else najlepsia
    c_odch, c_kroky, c_opt = 3, 3 + k_, 3 + 2 * k_
    c_suc = c_opt + 2
    sirky(wg, [14, 11] + [12] * (2 * k_ + 1) + [3] + [13] * (2 * k_ + 1))
    bunka(wg, "A1", "Výsledky podľa veľkosti úlohy", F_PODNADPIS)
    odsek(wg, 2, "Veľkosť = väčší z rozmerov tabuľky, max(m, n). Napr. úloha 5 × 12 má veľkosť 12. Pri väčších "
                 "úlohách je viac možností, ako prepravu rozvrhnúť, preto je ťažšie trafiť optimum. Stĺpce vpravo "
                 "(sivé) sú súčty, z ktorých sa počítajú priemery.", 9, F_SIVE)
    r0 = hlavicka(wg, 4, ["veľkosť max(m, n)", "počet úloh"] + [f"odchýlka {m}" for m in M]
                  + [f"kroky MODI {m}" for m in M] + [f"optimum {opt_m}"],
                  ["väčší rozmer tabuľky", "koľko úloh malo túto veľkosť"]
                  + ["priemerná odchýlka od optima"] * k_ + ["priemerný počet krokov MODI"] * k_
                  + [f"v koľkých % úloh {opt_m} trafila optimum hneď"])
    hlavicka(wg, 4, [f"súčet odchýlok {m}" for m in M] + [f"súčet krokov {m}" for m in M]
             + [f"optimum {opt_m} (počet)"], stlpec=c_suc)
    vel = sorted(s.velkost)
    for i, v in enumerate(vel):
        r = r0 + i
        g = s.velkost[v]
        bunka(wg, f"A{r}", v, F_TUCNE, okraj=True)
        bunka(wg, f"B{r}", g[M[0]][0], fmt=CELE, okraj=True)
        surove = [g[m][1] for m in M] + [g[m][2] for m in M] + [g[opt_m][3]]
        for j, h in enumerate(surove):
            bunka(wg, wg.cell(r, c_suc + j).coordinate, h, F_SIVE, "#,##0.00" if j < k_ else CELE, okraj=True)
        for j in range(k_):
            stl = get_column_letter(c_suc + j)
            bunka(wg, wg.cell(r, c_odch + j).coordinate, f"=IFERROR({stl}{r}/B{r}/100,0)", fmt=PCT, okraj=True)
            stl = get_column_letter(c_suc + k_ + j)
            bunka(wg, wg.cell(r, c_kroky + j).coordinate, f"=IFERROR({stl}{r}/B{r},0)", fmt=DES1, okraj=True)
        stl = get_column_letter(c_suc + 2 * k_)
        bunka(wg, wg.cell(r, c_opt).coordinate, f"=IFERROR({stl}{r}/B{r},0)", fmt=PCT, okraj=True)
    if vel:
        rk = r0 + len(vel) - 1
        for nazov, c1, c2, kotva, fmt in (("Priemerná odchýlka od optima podľa veľkosti", c_odch, c_kroky - 1, "A", "0%"),
                                          ("Priemerný počet krokov MODI podľa veľkosti", c_kroky, c_opt - 1,
                                           get_column_letter(max(10, c_opt + 1)), "0")):
            lg = LineChart()
            lg.title = nazov
            lg.y_axis.numFmt = fmt
            lg.add_data(Reference(wg, min_col=c1, max_col=c2, min_row=r0, max_row=rk), titles_from_data=False)
            lg.set_categories(Reference(wg, min_col=1, min_row=r0, max_row=rk))
            for ser, m in zip(lg.series, M):
                ser.tx = SeriesLabel(v=m)
                ser.graphicalProperties.line.solidFill = FARBY[m]
                ser.graphicalProperties.line.width = 28000
                ser.smooth = False
            osi(lg)
            lg.height, lg.width = 8, 14
            wg.add_chart(lg, f"{kotva}{rk + 3}")
    wg.freeze_panes = wg.cell(r0, 2)

    wt = wb.create_sheet("Podľa typu úlohy")
    sirky(wt, [20, 12] + [13] * (3 * k_))
    bunka(wt, "A1", "Výsledky podľa typu úlohy", F_PODNADPIS)
    odsek(wt, 2, "Vyrovnaná úloha: dodávatelia majú spolu presne toľko, koľko odberatelia chcú. Prebytok ponuky: "
                 "dodávatelia majú viac → pridá sa fiktívny odberateľ, ktorý „zoberie“ zvyšok za cenu 0. Prebytok "
                 "dopytu: odberatelia chcú viac → pridá sa fiktívny dodávateľ s cenou 0.", 11, F_SIVE)
    r0 = hlavicka(wt, 4, ["typ úlohy", "počet úloh"] + [f"odchýlka {m}" for m in M]
                  + [f"kroky MODI {m}" for m in M] + [f"optimum {m}" for m in M],
                  ["", ""] + ["priemerná odchýlka od optima"] * k_ + ["priemerný počet krokov MODI"] * k_
                  + ["trafila optimum hneď"] * k_)
    for i, (t_, g) in enumerate(sorted(s.typ.items(), key=lambda x: list(TYPY).index(x[0]) if x[0] in TYPY else 9)):
        r = r0 + i
        n_ = max(1, g[M[0]][0])
        bunka(wt, f"A{r}", TYPY.get(t_, t_), F_TUCNE, okraj=True)
        bunka(wt, f"B{r}", g[M[0]][0], fmt=CELE, okraj=True)
        for j, m in enumerate(M):
            bunka(wt, wt.cell(r, 3 + j).coordinate, g[m][1] / n_ / 100, fmt=PCT, okraj=True)
            bunka(wt, wt.cell(r, 3 + k_ + j).coordinate, g[m][2] / n_, fmt=DES1, okraj=True)
            bunka(wt, wt.cell(r, 3 + 2 * k_ + j).coordinate, g[m][3] / n_, fmt=PCT2, okraj=True)
    bunka(wt, f"A{r0 + len(s.typ) + 1}", "Hodnoty sú priemery vypočítané programom zo súčtov počas simulácie.", F_SIVE)

    wh = wb.create_sheet("Rozdelenie odchýlky")
    sirky(wh, [16] + [12] * (2 * k_))
    bunka(wh, "A1", "Rozdelenie odchýlky od optima", F_PODNADPIS)
    odsek(wh, 2, "Koľko úloh padlo do ktorého pásma odchýlky. Napr. riadok „0–5 %“ = úlohy, kde bolo prvé riešenie "
                 "najviac o 5 % drahšie ako optimum (vrátane úloh, kde ho metóda trafila presne). Posledný riadok "
                 "zahŕňa všetko od 495 % vyššie. Čím viac úloh je v horných riadkoch, tým je metóda lepšia.", 7, F_SIVE)
    r0 = hlavicka(wh, 4, ["pásmo odchýlky"] + [f"počet úloh {m}" for m in M] + [f"podiel {m}" for m in M])
    hrany = HIST_HRANY
    for i in range(len(hrany) - 1):
        r = r0 + i
        text = f"{hrany[i]:.0f}–{hrany[i + 1]:.0f} %" if i < len(hrany) - 2 else f"{hrany[i]:.0f} % a viac"
        bunka(wh, f"A{r}", text, okraj=True)
        for j, m in enumerate(M):
            bunka(wh, wh.cell(r, 2 + j).coordinate, int(s.hist[m][i]), fmt=CELE, okraj=True)
            stl = get_column_letter(2 + j)
            bunka(wh, wh.cell(r, 2 + k_ + j).coordinate,
                  f"=IFERROR({stl}{r}/SUM({stl}${r0}:{stl}${r0 + len(hrany) - 2}),0)", fmt=PCT, okraj=True)
    rk = r0 + len(hrany) - 2
    lg = LineChart()
    lg.title = "Podiel úloh podľa odchýlky od optima"
    lg.y_axis.scaling.min = 0
    lg.y_axis.numFmt = "0%"
    lg.add_data(Reference(wh, min_col=2 + k_, max_col=1 + 2 * k_, min_row=r0 - 1, max_row=rk), titles_from_data=True)
    lg.set_categories(Reference(wh, min_col=1, min_row=r0, max_row=rk))
    for ser, m in zip(lg.series, M):
        ser.graphicalProperties.line.solidFill = FARBY[m]
        ser.graphicalProperties.line.width = 22000
    osi(lg)
    lg.height, lg.width = 9, 18
    wh.add_chart(lg, f"{get_column_letter(2 * k_ + 3)}4")
    wh.freeze_panes = wh.cell(r0, 2)

    wu = wb.create_sheet("Údaje (vzorka)")
    sirky(wu, [10, 9, 9, 16, 10, 14, 9, 12, 12, 12, 12, 12, 12, 12, 12])
    pocet_vz = len(s.vzorka) // k_
    odsek(wu, 1, f"Každý riadok = jedna úloha vyriešená jednou metódou (každá úloha má {k_} riadky: {zoznam}). "
                 f"Je tu prvých {cislo(pocet_vz, 0)} úloh z {cislo(s.hotovo, 0)}. Šípkami v hlavičke sa dá "
                 f"filtrovať (napr. len VAM alebo len úlohy 3 × 3).", 15, F_SIVE)
    nazvy_u = ["číslo úlohy", "dodávatelia m", "odberatelia n", "typ úlohy", "ceny 1 až", "fiktívny účastník",
               "metóda", "doplnené nuly v báze", "z (prvé riešenie)", "z* (optimum)", "odchýlka od optima",
               "trafila optimum", "čas metódy (µs)", "čas MODI (µs)", "kroky MODI"]
    vysv_u = ["poradové číslo (spolu so seedom ju presne určuje)", "počet riadkov", "počet stĺpcov",
              "vyrovnaná / prebytok ponuky / dopytu", "najvyššia možná cena prepravy",
              "pridaný riadok alebo stĺpec s cenou 0", "skratka metódy",
              "koľko polí sa doplnilo nulou (degenerácia)", "celková cena prvého riešenia",
              "najnižšia možná celková cena", "(z − z*) / z*", "áno = prvé riešenie je hneď optimálne",
              "mikrosekundy (1 µs = 0,001 ms)", "mikrosekundy", "počet opráv metódou MODI"]
    r = hlavicka(wu, 2, nazvy_u, vysv_u)
    zaciatok = r
    for riadok in s.vzorka:
        (id_, m_, n_, typ_, cena, fikt, met, deg, z, zopt, odch, opt, tmet, tmodi, kroky) = riadok
        wu.append([id_, m_, n_, TYPY.get(typ_, typ_), cena,
                   {"dodavatel": "dodávateľ", "odberatel": "odberateľ"}.get(fikt, fikt or "–"),
                   met, deg, z, zopt, odch / 100, "áno" if opt else "nie", round(tmet / 1e3, 1),
                   round(tmodi / 1e3, 1), kroky])
    for stl, fmt in (("I", CELE), ("J", CELE), ("K", PCT), ("M", DES1), ("N", DES1)):
        for c in wu[f"{stl}{zaciatok}:{stl}{wu.max_row}"]:
            c[0].number_format = fmt
    wu.freeze_panes = wu.cell(zaciatok, 1)
    wu.auto_filter.ref = f"A{zaciatok - 2}:O{wu.max_row}"

    wsl = wb.create_sheet("Slovník pojmov")
    sirky(wsl, [30, 100])
    bunka(wsl, "A1", "Slovník pojmov – čo znamenajú skratky", F_PODNADPIS)
    r = hlavicka(wsl, 3, ["pojem", "vysvetlenie (ľudskou rečou)"])
    pojmy = [
        ("Dopravná úloha", "Máme dodávateľov s určitým množstvom tovaru a odberateľov, ktorí tovar potrebujú. Poznáme "
                           "cenu prepravy jednej jednotky z každého dodávateľa ku každému odberateľovi. Hľadáme plán "
                           "prepravy (kto komu koľko pošle), pri ktorom sú celkové náklady najnižšie."),
        ("m, n", "m = počet dodávateľov (riadky tabuľky), n = počet odberateľov (stĺpce tabuľky)."),
        ("c_ij", "Cena prepravy jednej jednotky od dodávateľa i k odberateľovi j."),
        ("a_i, b_j", "a_i = koľko tovaru má dodávateľ i (kapacita), b_j = koľko tovaru chce odberateľ j (požiadavka)."),
        ("z", "Celková cena prepravného plánu = súčet (prepravené množstvo × cena) cez všetky polia tabuľky."),
        ("Východiskové riešenie", "Prvý prípustný plán prepravy (všetko sa rozvezie, nikto nedostane viac/menej). "
                                  "Nemusí byť najlacnejší – slúži ako štart pre hľadanie optima."),
        ("Optimum, z*", "Najlacnejší možný plán prepravy a jeho cena. Našli sme ho metódou MODI."),
        *[(f"{m} – {NAZVY[m].lower()}", POJMY_METOD[m]) for m in M],
        ("MODI (metóda potenciálov)", "Postup, ktorý z východiskového riešenia krok za krokom nájde optimum. Spočíta "
                                      "tzv. potenciály u_i a v_j a zistí, či sa dá plán ešte zlacniť. Ak áno, presunie "
                                      "prepravu po uzavretom obvode v tabuľke (jeden krok MODI) a opakuje to."),
        ("Krok MODI", "Jedno zlepšenie plánu presunom prepravy po uzavretom obvode. Menej krokov = menej práce "
                      "= východiskové riešenie bolo bližšie k optimu."),
        ("Odchýlka od optima (%)", "(z − z*) / z* × 100. Hovorí, o koľko percent je prvé riešenie drahšie ako "
                                   "najlacnejšie možné. 0 % = metóda trafila optimum hneď."),
        ("Trafila optimum", "Prvé riešenie metódy bolo hneď najlacnejšie možné (odchýlka 0 %)."),
        ("Najlepšia v úlohe", "V tej istej úlohe dala metóda najlacnejšie zo všetkých prvých riešení. Ak dve metódy "
                              "dali rovnako dobré riešenie, počíta sa obom."),
        ("Vyrovnaná / nevyrovnaná úloha", "Vyrovnaná: celková ponuka = celkový dopyt. Nevyrovnaná: pridá sa "
                                          "fiktívny dodávateľ alebo odberateľ s cenou 0, aby sa vyrovnala."),
        ("Degenerácia", "Plán má mať presne m + n − 1 obsadených polí. Ak ich vyjde menej, doplní sa pole s "
                        "prepravou 0, aby sa dala použiť metóda MODI."),
        ("Seed", "Číslo, z ktorého počítač vyrába náhodné úlohy. Rovnaký seed = rovnaké úlohy, výsledky sa dajú "
                 "zopakovať a skontrolovať."),
        ("Priemer", "Súčet hodnôt deleno ich počtom."),
        ("ms, µs", "Milisekunda = tisícina sekundy, mikrosekunda = milióntina sekundy."),
    ]
    for pojem, text in pojmy:
        bunka(wsl, f"A{r}", pojem, F_TUCNE, align=ZALOM, okraj=True)
        bunka(wsl, f"B{r}", text, align=ZALOM, okraj=True)
        wsl.row_dimensions[r].height = 14 * (len(text) // 110 + 1) + 4
        r += 1
    r += 1
    bunka(wsl, f"A{r}", "Ako to vysvetliť vyučujúcemu", F_PODNADPIS)
    r += 1
    b = najlepsia
    vysvetlenie = [
        f"1. Cieľ: zistiť, ktorá z metód {zoznam} dáva najlepšie východiskové riešenie a koľko "
        "práce potom zostane metóde MODI.",
        f"2. Postup: program vygeneroval {cislo(s.hotovo, 0)} náhodných úloh rôznych veľkostí, s rôznymi cenami, "
        "vyrovnané aj nevyrovnané. Každú vyriešil všetkými porovnávanými metódami a potom MODI do optima.",
        "3. Správnosť programu: metódy sme overili na príkladoch z prednášky (výsledky sedia) a optimum sme "
        "porovnali s nezávislým riešičom (knižnica scipy) – zhoda vo všetkých kontrolovaných úlohách.",
        "4. Meranie: pre každú úlohu a metódu sme zapísali cenu prvého riešenia z, optimum z*, odchýlku "
        "(z − z*)/z*, počet krokov MODI a čas.",
        f"5. Výsledok: najmenšiu odchýlku aj najmenej krokov MODI má {b} (pozri hárok „Výsledky metód“). "
        "Rozdiel medzi metódami rastie s veľkosťou úlohy (hárok „Podľa veľkosti“).",
        "6. Zopakovateľnosť: úlohy sú určené seedom, takže ten istý výpočet sa dá kedykoľvek zopakovať s rovnakým "
        "výsledkom.",
    ]
    for veta in vysvetlenie:
        wsl.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        bunka(wsl, f"A{r}", veta, align=ZALOM)
        wsl.row_dimensions[r].height = 30
        r += 1

    poradie = ["Prehľad", "Výsledky metód", "Podľa veľkosti", "Podľa typu úlohy", "Rozdelenie odchýlky",
               "Údaje (vzorka)", "Slovník pojmov", "Podklady"]
    wb._sheets = [wb[n] for n in poradie]
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    wb["Údaje (vzorka)"].sheet_view.showGridLines = True
    wb.calculation.fullCalcOnLoad = True
    wb.properties.title = "Porovnanie metód dopravnej úlohy"
    cesta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cesta)
    return cesta


def info_z_json(cesta_json: Path) -> dict:
    d = json.loads(Path(cesta_json).read_text(encoding="utf-8"))
    return {
        "nazov": d.get("nazov"), "seed": d.get("seed"), "procesy": d.get("procesy"),
        "pocet_zadany": d.get("pocet_uloh_zadany", d.get("pocet_uloh")),
        "parametre_generatora": d.get("parametre_generatora", {}),
        "trvanie_s": d.get("trvanie_behu_s", 0), "datum": d.get("datum", ""),
    }


def main():
    if len(sys.argv) < 2:
        print("Použitie: .venv\\Scripts\\python.exe export_excel.py <názov behu>   (napr. hlavny100k)")
        sys.exit(1)
    nazov = sys.argv[1]
    cesta_csv = KOREN / "vysledky" / f"beh_{nazov}.csv"
    t0 = time.perf_counter()
    s = Statistiky.z_csv(cesta_csv)
    info = info_z_json(cesta_csv.with_suffix(".json"))
    vystup = exportuj(KOREN / "vysledky" / f"{nazov}.xlsx", s, info)
    print(f"Uložené: {vystup}  ({time.perf_counter() - t0:.1f} s)")


if __name__ == "__main__":
    main()
